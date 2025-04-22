import os
import openpyxl
import sqlparse
import pandas as pd
from psycopg2 import sql 
from dotenv import load_dotenv
from db_connection import PostgresDB

# Map excel sheet names to SQL table names
TABLE_MAP = {
        'Client'       : 'client',
        'Transaction'  : 'transaction',
        'MediaItem'    : 'media_item',
        'Book'         : 'book',
        'Magazine'     : 'magazine',
        'DigitalMedia' : 'digital_media',
    }

def main(test=False, build_tables=False, drop_tables=False):
    """Execute the data parsing and population logic."""
    load_dotenv()

    if drop_tables:
        conn, _ = open_db_conn()
        if not conn:
            raise Exception("Failed to establish database connection for dropping tables.")
        cursor = conn.cursor()
        try:
            drop_table(cursor, conn)
            print("[SUCCESS] All tables dropped successfully.")
        finally:
            cursor.close()
            conn.close()
            if not build_tables:  # Exit if we're only dropping tables
                return

    # If informed to do so, create the tables to inhabit the database
    if build_tables: 
        # Relative path to the ddl file
        ddl_path = "src/libraryDDL.sql"
        if not (create_tables(ddl_path)): 
            raise Exception("Failed to create tables from DDL file.")

    PATH = os.getenv("EXCEL_PATH")
    
    SHEET_NAMES = ['Client', 'Transaction', 'MediaItem', 'Book', 'Magazine', 'DigitalMedia']



    if not PATH: 
        raise ValueError("EXCEL_PATH is not set in the environment.")
    
    tables_data = {}

    # Parse data    
    for sheet in SHEET_NAMES: 
        workbook = read_sheet(PATH, sheet)  
        tables_data[sheet] = parse_workbook(sheet, workbook)            

    # Populate DB
    for sheet in SHEET_NAMES:     
        data = tables_data[sheet]
        if not data: 
            print(f"[SKIPPED] {sheet}: empty or failed to parse. No rows inserted.")
            continue
            
        # Include test routing. Shows data to be inserted to the db, in the correct order
        if test: 
            target_table = "Book"
            populate_table_test(sheet, data, target_table)
        else:        
            populate_table(sheet, data) 

def open_db_conn(): 
    """Establish a connection to the PostgreSQL database."""
    try: 
        db = PostgresDB(password=os.getenv("DB_PASSWORD"))
        conn = db.connect()
        return conn, db
    except Exception as e:
        print(f"[ERROR] Failed to open DB connection: {e}")
        return None, None
    
def create_tables(ddl_path) -> bool:
    """Create database relations by executing the DDL statements."""
    # Open a connection to the database
    conn, _ = open_db_conn()

    if not conn:
        print("[ERROR] Could not establish a database connection.")
        return

    try: 
        # Read the .sql file at the path
        with open(ddl_path, 'r') as ddl_file:
            # Store the statements in the file
            ddl_content = ddl_file.read()

        # Create a cursor object
        cursor = conn.cursor()

        # Use sqlparse to split the the statements
        # - Sqlparse parses the sql into tokens and understands the structure of the statements 
        # - Better than a split by semicolon, for example
        statements = sqlparse.split(ddl_content)

        success = True

        # Execute each statement
        for statement in statements:
            statement = statement.strip() 
            if statement:
                try:
                    cursor.execute(statement)
                    print("[SUCCESS] Executed statement")
                except Exception as e:
                    success = False # Mark failure
                    print(f"[ERROR] Failed to execute statement: {e}")
        
        conn.commit() # Commit changes
        print("[SUCCESS] Database tables successfully created.\n")
    except Exception as e: 
        print(f"[ERROR] Failed to execute DDL statements: {e}")

    # Irrespective of the result, terminate connection and clean up
    finally: 
        if cursor: 
            cursor.close()
        if conn: 
            conn.close()
    
    return success

def read_sheet(file_path, excel_sheet_name=None): 
    """Read a specified sheet in a given excel file."""
    try:
        return pd.read_excel(file_path, sheet_name=excel_sheet_name)
    except Exception as e: 
        print(f"Failed to read sheet '{excel_sheet_name}': '{e}'")
        return pd.DataFrame() # Return empty dataframe to safely skip

def parse_workbook(table_name, workbook):
    """Parse a given workbook, store information."""
    sheet_store = []

    # Store the information as a list of dict objects
    for _, row in workbook.iterrows(): 
        # Skip completely empty rows
        if row.isnull().all():
            continue

        # Skip rows with only occupied cell
        for index, cell in enumerate(row):
            if index == 0:
                continue
            if not pd.isnull(cell):
                break
        else:
            continue

        data = {}

        for col_name in workbook.columns:
            value = row[col_name]

            if pd.isna(value):
                value = None
            elif isinstance(value, float) and value.is_integer():
                value = int(value)
            elif isinstance(value, pd.Timestamp):
                if pd.isna(value):
                    value = None
                elif value.year < 1970: # Check for excel corruption issue (1969-12-31 18:00:00). Was triggered by cell arithmetic
                    value = None
                    print(f"[WARNING] Skipping invalid date in column {col_name}. Row: {row}")
                else: 
                    value = value.to_pydatetime()

            data[col_name.lower()] = value

        # ISBN sanitization
        if 'isbn' in data and data['isbn'] is not None: 
            data['isbn'] = str(data['isbn']).strip().replace('.0', '')

        sheet_store.append(data)
        
    return sheet_store

def populate_table_test(table_name, table_data, target_table, verbose=True): 
    if not target_table.strip(): raise ValueError("Invalid target table name.")
    if table_name == target_table:
        for row in table_data:
            print(f"[INSERTED into {table_name}] {row}\n")
      

def populate_table(sheet_name, table_data): 
    """Populate a table in the PostgreSQL database."""
    conn, db = open_db_conn()

    if not conn:
        print(f"[SKIPPED] {sheet_name}: could not connect to DB.")
        return

    cursor = conn.cursor()
    
    for row in table_data: 
        try: 
            insert_row(cursor, sheet_name, row)
            conn.commit() # Commit this row insertion
            print(f"[INSERTED into {TABLE_MAP[sheet_name]}] {row}")
        except Exception as e: 
            conn.rollback() # Undo the failed insert
            print(f"[ERROR] Skipping row in {TABLE_MAP[sheet_name]}: {e}")
    
    cursor.close()
    db.close()

def insert_row(cursor, sheet_name, row_data): 
    """Insert a row into a table in the PostgreSQL database."""
    table_name = TABLE_MAP[sheet_name]
    
    # Column name mapping for Transaction table
    COLUMN_MAP = {
        'date_returned': 'returned_date'
    }

    # Since the transaction ID is a SERIAL PRIMARY KEY in the ddl, remove manually set ID fields
    if sheet_name == 'Transaction': 
        transaction_key = next((k for k in row_data if k.lower() == 'transaction_id'), None)
        if transaction_key: 
            row_data.pop(transaction_key)

    # For the derived tables (B, M, DM), make sure the item id exists (since it'll be linked to Media Item using it
    if sheet_name in ['Book', 'Magazine', 'DigitalMedia']: 
        item_id = row_data.get('item_id')
        if item_id is None: 
            raise ValueError(f"{sheet_name} row missing item_id, which is required to link to Media_Item")
    
    # Map column names if needed
    mapped_data = {}
    for k, v in row_data.items():
        mapped_key = COLUMN_MAP.get(k, k)  # Use mapped name if it exists, otherwise use original
        mapped_data[mapped_key] = v

    columns = list(mapped_data.keys())
    values = list(mapped_data.values())

    # Use psycopg2.sql to safely interpolate identifiers and placeholders
    query = sql.SQL("INSERT INTO {table} ({fields}) VALUES ({placeholders})").format(
            table = sql.Identifier(table_name),
            fields=sql.SQL(', ').join(
                sql.Identifier(str(c).lower()) for c in columns
            ),
            placeholders=sql.SQL(', ').join(sql.Placeholder() * len(columns))
    )

    # Any error here will bubble up into the populate_table() catch block
    cursor.execute(query, values)

def drop_table(cursor, conn): 
    cursor.execute("""
        DROP TABLE IF EXISTS "Transaction" CASCADE;
        DROP TABLE IF EXISTS book           CASCADE;
        DROP TABLE IF EXISTS magazine       CASCADE;
        DROP TABLE IF EXISTS digital_media  CASCADE;
        DROP TABLE IF EXISTS media_item     CASCADE;
        DROP TABLE IF EXISTS client         CASCADE;

        DROP TYPE  IF EXISTS availability_status_enum  CASCADE;
        DROP TYPE  IF EXISTS account_status_enum       CASCADE;
        DROP TYPE  IF EXISTS membership_type_enum      CASCADE;
        """)
    
    conn.commit()


if __name__ == "__main__":
    # To drop and rebuild tables:
    main(drop_tables=True, build_tables=True)
    
    # To only drop tables:
    # main(drop_tables=True)
    
    # To only build tables:
    # main(build_tables=True)